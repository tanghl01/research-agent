#!/usr/bin/env python3
"""Small PPTX inspection helper used by @operator/core validation checks."""

from __future__ import annotations

import json
import sys
from collections.abc import Iterable
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main(argv: list[str]) -> int:
    if len(argv) < 3:
        print(
            "Usage: pptx_inspect.py <slide-count|text-coverage|image-count> <pptx> [payload_json]",
            file=sys.stderr,
        )
        return 2

    command = argv[1]
    pptx_path = argv[2]
    presentation = Presentation(pptx_path)

    if command == "slide-count":
        print(len(presentation.slides))
        return 0

    if command == "text-coverage":
        if len(argv) < 4:
            print("text-coverage requires a JSON payload", file=sys.stderr)
            return 2
        report = text_coverage(presentation, json.loads(argv[3]))
        print(json.dumps(report, ensure_ascii=False))
        return 0

    if command == "image-count":
        print(sum(count_pictures(slide.shapes) for slide in presentation.slides))
        return 0

    print(f"Unknown command: {command}", file=sys.stderr)
    return 2


def text_coverage(presentation: Any, config: dict[str, Any]) -> dict[str, Any]:
    text = "\n".join(
        shape.text for slide in presentation.slides for shape in walk_shapes(slide.shapes) if has_text(shape)
    )
    required = [term for term in config.get("required_terms", []) if term]
    forbidden = [term for term in config.get("forbidden_terms", []) if term]
    matched = [term for term in required if term in text]
    return {
        "matched": matched,
        "missing": [term for term in required if term not in text],
        "forbidden_hits": [term for term in forbidden if term in text],
        "text_chars": len(text),
    }


def count_pictures(shapes: Iterable[Any]) -> int:
    total = 0
    for shape in walk_shapes(shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            total += 1
    return total


def walk_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        nested_shapes = getattr(shape, "shapes", None)
        if nested_shapes is not None:
            yield from walk_shapes(nested_shapes)


def has_text(shape: Any) -> bool:
    return bool(getattr(shape, "has_text_frame", False))


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
